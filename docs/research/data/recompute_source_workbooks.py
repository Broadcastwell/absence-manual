"""Read original deposited workbooks without recalculating or saving them.

Run with an input directory containing Zenodo 3381151 and 3484673 and an
output directory. Cached source values and reconstructed results stay distinct.
Dependencies: openpyxl, numpy, pandas. No network requests.
"""
from pathlib import Path
import argparse, csv, hashlib, json, math, statistics
import numpy as np
import pandas as pd
import openpyxl

def rho(a, b):
    ar = pd.Series(a).rank(method='average').to_numpy()
    br = pd.Series(b).rank(method='average').to_numpy()
    return float(np.corrcoef(ar, br)[0, 1])

def dump(path, rows):
    with path.open('w', encoding='utf-8', newline='') as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader(); w.writerows(rows)

def main(source, out):
    out.mkdir(parents=True, exist_ok=True)
    configurations = []; differences = []
    for p in sorted((source/'3381151').glob('*.xlsx')):
        wb = openpyxl.load_workbook(p, data_only=True)
        ws = wb.worksheets[0]
        labels = [c.value for c in ws[1]]
        median_col = next(i+1 for i, x in enumerate(labels)
                          if str(x).lower() in ('median', 'mediana'))
        query_cols = [i+1 for i, x in enumerate(labels)
                      if i > 0 and isinstance(x, str)
                      and x.lower() not in ('median', 'mediana')]
        assert len(query_cols) == 25
        rank = []; cached = []; full = []; available = []
        for row in range(2, 1002):
            rank.append(ws.cell(row, 1).value)
            cached.append(ws.cell(row, median_col).value)
            values = [ws.cell(row, c).value for c in query_cols]
            assert all(x is None or isinstance(x, (int, float)) for x in values)
            numeric=[x for x in values if isinstance(x, (int, float))]
            available.append(len(numeric))
            full.append(statistics.median(numeric))
            if cached[-1] != full[-1]:
                differences.append(dict(file=p.name, sheet=ws.title,
                    row=row, cached_median=cached[-1],
                    recomputed_25_query_median=full[-1]))
        formulas = openpyxl.load_workbook(p, data_only=False).worksheets[0]
        configurations.append(dict(file=p.name, sheet=ws.title,
            sha256=hashlib.sha256(p.read_bytes()).hexdigest(),
            query_columns=25, result_positions=len(rank),
            blank_query_cells=25000-sum(available),
            minimum_available_queries_at_a_position=min(available),
            first_median_formula=formulas.cell(2, median_col).value,
            cached_rho=rho(rank, cached),
            full_25_query_rho=rho(rank, full),
            differing_medians=sum(x != y for x,y in zip(cached,full))))
        wb.close()
    dump(out/'rovira-workbook-reanalysis.csv', configurations)
    dump(out/'rovira-median-differences.csv', differences)

    p=source/'3484673/Results - Academic Excellence and Web Presence Correlation.xlsx'
    w=openpyxl.load_workbook(p,data_only=True); s=w.worksheets[0]
    data=[list(row) for row in s.iter_rows(min_row=3,max_row=102,values_only=True)]
    assert len(data)==100 and len(set(x[0] for x in data))==100
    rank_rho=rho([x[2] for x in data],[x[7] for x in data])
    # Conventional two-sided t approximation for the rank correlation.
    df=98; t=abs(rank_rho)*math.sqrt(df/(1-rank_rho**2))
    # Simpson integration of the Student-t density from 0 to |t|.
    xs=np.linspace(0,t,10001)
    coeff=math.exp(math.lgamma((df+1)/2)-math.lgamma(df/2))/math.sqrt(df*math.pi)
    ys=coeff*(1+xs*xs/df)**(-(df+1)/2)
    area=(t/10000)/3*(ys[0]+ys[-1]+4*sum(ys[1:-1:2])+2*sum(ys[2:-1:2]))
    r2=float(np.corrcoef([x[1] for x in data],[x[3] for x in data])[0,1]**2)
    summary=dict(source_sha256=hashlib.sha256(p.read_bytes()).hexdigest(),
        sheet=s.title,rows='3:102',arwu_rank_column='C',seo_rank_column='H',
        n=100,arwu_seo_rank_spearman=rank_rho,
        two_sided_t_approximation_p=1-2*area,
        arwu_rating_quality_rating_r_squared=r2,
        arwu_rating_quality_rating_F_1_98=r2/(1-r2)*98,
        qualification='Deposited composite scores and ranks, not new website audits or observed search positions.')
    (out/'university-workbook-reanalysis.json').write_text(json.dumps(summary,indent=2),encoding='utf-8')
    dump(out/'university-source-results.csv',[dict(zip(
        ['institution','arwu_rating','arwu_rank','quality_rating','seo_rating','overall_rating','quality_rank','seo_rank','overall_rank'],x)) for x in data])
    print(json.dumps({'rovira_configurations':configurations,'university':summary},indent=2))

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--source',type=Path,required=True);p.add_argument('--out',type=Path,required=True)
    args=p.parse_args();main(args.source,args.out)
